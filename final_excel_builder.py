import openpyxl
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image
from io import BytesIO
import os
import pandas as pd

def create_final_excel(scraped_excel_path, final_excel_path, data_excel_path=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Product Data"

    scraped_wb = openpyxl.load_workbook(scraped_excel_path)
    scraped_ws = scraped_wb.active

    header_row = [cell.value.strip() if cell.value else "" for cell in next(scraped_ws.iter_rows(min_row=1, max_row=1))]

    # Standardized column mapping
    col_map = {
        'id_col': 'Model Number' if 'Model Number' in header_row else 'Part Number',
        'price_col': 'Price' if 'Price' in header_row else 'Unit Price',
        'brand_col': 'Brand',
        'qty_col': 'Qty' if 'Qty' in header_row else 'Quantity',
        'desc_col': 'Description',
        'specs_col': 'Specifications',
        'img_col': 'Image URL'
    }

    # Create main Excel with embedded images
    ws.append([
        col_map['id_col'],
        col_map['desc_col'],
        col_map['price_col'],
        col_map['specs_col'],
        col_map['img_col'],
        "Embedded Image"
    ])

    row_num = 2

    for row in scraped_ws.iter_rows(min_row=2, values_only=True):
        identifier = str(row[header_row.index(col_map['id_col'])]) if row[header_row.index(col_map['id_col'])] else ""
        
        if not identifier:
            continue

        desc = str(row[header_row.index(col_map['desc_col'])] if col_map['desc_col'] in header_row else "").strip()
        price = str(row[header_row.index(col_map['price_col'])] if col_map['price_col'] in header_row else "").strip()
        specs = str(row[header_row.index(col_map['specs_col'])] if col_map['specs_col'] in header_row else "").strip()
        img_url = str(row[header_row.index(col_map['img_col'])] if col_map['img_col'] in header_row else "").strip()

        ws.append([identifier, desc[:300], price or "N/A", specs[:300] or "N/A", img_url or "", ""])

        # Look for image with original identifier (spaces preserved)
        img_path = None
        for ext in [".jpg", ".jpeg", ".png"]:
            candidate = os.path.join("downloaded_images", f"{identifier}{ext}")
            if os.path.exists(candidate):
                img_path = candidate
                break

        if img_path:
            try:
                with Image.open(img_path) as im:
                    im.thumbnail((100, 100))
                    buffer = BytesIO()
                    im.save(buffer, format="PNG")
                    buffer.seek(0)

                    img = ExcelImage(buffer)
                    img.width, img.height = 80, 80
                    ws.add_image(img, f"F{row_num}")
                    ws.row_dimensions[row_num].height = 65
            except Exception as e:
                print(f"❌ Failed to embed image for {identifier}: {e}")

        row_num += 1

    # Set column widths
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 15
    ws.column_dimensions["D"].width = 60
    ws.column_dimensions["E"].width = 50
    ws.column_dimensions["F"].width = 18

    wb.save(final_excel_path)

    # Create simplified data Excel for frontend
    if data_excel_path:
        df = pd.read_excel(scraped_excel_path)
        simplified_df = pd.DataFrame()
        
        simplified_df['Model Number'] = df[col_map['id_col']]
        simplified_df['Price'] = df.get(col_map['price_col'], 'N/A')
        simplified_df['Brand'] = df.get(col_map['brand_col'], 'Unknown')
        simplified_df['Qty'] = df.get(col_map['qty_col'], 'N/A')
        simplified_df['Description'] = df.get(col_map['desc_col'], '')
        simplified_df['Specifications'] = df.get(col_map['specs_col'], '')
        
        if col_map['img_col'] in df.columns:
            simplified_df['Image URL'] = df[col_map['img_col']]
        
        simplified_df.to_excel(data_excel_path, index=False)